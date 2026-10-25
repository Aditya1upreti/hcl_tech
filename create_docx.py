from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_docx(filename, title, content_lines):
    doc = Document()
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    
    # Add title
    title_para = doc.add_heading(title, level=1)
    title_para.runs[0].font.color.rgb = RGBColor(0, 51, 102)
    
    # Add content
    for line in content_lines:
        if line.startswith('=') or line.startswith('#'):
            heading_level = 2 if line.startswith('==') else 1
            para = doc.add_heading(line.strip('=').strip('#'), level=heading_level)
            para.runs[0].font.color.rgb = RGBColor(0, 51, 102)
        elif line.strip():
            para = doc.add_paragraph(line.strip())
            # Adjust spacing
            para.paragraph_format.space_after = Pt(6)
    
    doc.save(filename)
    print(f"Created {filename}")

# Leave Policy DOCX
leave_content = [
    "LEAVE POLICY",
    "HCLTech Career Shaper Programme",
    "",
    "1. PURPOSE",
    "This policy establishes the leave entitlements and procedures for all participants in the HCLTech Career Shaper Programme.",
    "",
    "2. LEAVE ENTITLEMENTS",
    "- Casual Leave: 10 days per year (maximum 3 consecutive days)",
    "- Sick Leave: 15 days per year (certificate required for absences exceeding 3 days)",
    "- Medical Leave: As per company policy with proper documentation",
    "- Exam Leave: 5 days per examination period with prior approval",
    "- Parental Leave: As per company policy and applicable law",
    "",
    "3. LEAVE APPLICATION PROCEDURE",
    "- All leave must be applied for through the leave management system at least 5 working days in advance",
    "- Emergency leave may be applied for retroactively with supervisor approval",
    "- Last-minute leave (same day) is permitted for genuine emergencies only",
    "",
    "4. APPROVAL PROCESS",
    "- Line managers approve leave requests within their team",
    "- Department heads approve leave exceeding 5 consecutive days",
    "- The programme coordinator has final approval authority for all leave",
    "",
    "5. LEAVE MONITORING",
    "- Leave balances are updated monthly in the HR system",
    "- Excessive leave usage (more than 50% of entitlement) will be flagged for review",
    "- Unused casual leave may be carried forward up to 5 days at year-end",
    "",
    "6. SPECIAL LEAVE CATEGORIES",
    "- Bereavement Leave: 3 days with proper documentation",
    "- Compassionate Leave: Up to 7 days as per circumstances",
    "- Deputed Training Leave: For approved external training programmes",
    "",
    "7. NON-COMPLIANCE",
    "- Leave applied for without proper reason may result in denial",
    "- Unauthorized absence is treated as a conduct matter",
    "- Repeated policy violations may affect future leave entitlements",
    "",
    "8. REVIEW",
    "- This policy will be reviewed annually",
    "- Suggestions for improvement may be submitted to the programme coordinator",
]
create_docx("data/raw/leave_policy.docx", "Leave Policy", leave_content)

# Code of Conduct DOCX
conduct_content = [
    "CODE OF CONDUCT",
    "HCLTech Career Shaper Programme",
    "",
    "1. PURPOSE",
    "This code of conduct establishes the expected standards of behaviour for all participants in the HCLTech Career Shaper Programme. It aims to foster a respectful, inclusive, and productive learning environment.",
    "",
    "2. RESPECT AND INCLUSION",
    "- All participants must treat each other with respect and dignity",
    "- Discrimination, harassment, or bullying of any kind will not be tolerated",
    "- Participants must respect diversity in all its forms (cultural, ethnic, gender, religious, ability, etc.)",
    "- Inclusive language must be used in all communications",
    "",
    "3. PROFESSIONALISM",
    "- Participants must maintain professional conduct at all times",
    "- Punctuality is expected for all scheduled sessions and meetings",
    "- Deadlines must be met and commitments honoured",
    "",
    "4. CONFIDENTIALITY",
    "- All programme materials, assessments, and participant information must be kept confidential",
    "- Unauthorized disclosure of confidential information is prohibited",
    "- Participants must sign the non-disclosure agreement provided",
    "",
    "5. INTELLECTUAL PROPERTY",
    "- All intellectual property created during the programme remains the property of HCLTech",
    "- Participants must not claim ownership of programme materials or facilitator content",
    "- Open source contributions must acknowledge the programme as the source",
    "",
    "6. TECHNOLOGY USE",
    "- Programme systems and resources must be used for programme-related purposes only",
    "- Unauthorized downloading, installation, or use of software is prohibited",
    "- Personal use of programme resources must be reasonable and not interfere with programme activities",
    "",
    "7. ACADEMIC INTEGRITY",
    "- All assessments must be completed honestly and independently unless collaboration is explicitly permitted",
    "- Plagiarism in any form will not be tolerated",
    "- Cheating on examinations or assignments will result in disciplinary action",
    "- Participants must acknowledge all sources and references properly",
    "",
    "8. REPORTING VIOLATIONS",
    "- Participants must report any observed violations of this code through the established channel",
    "- Retaliation against anyone reporting a violation is prohibited",
    "- Reports will be investigated promptly and confidentially",
    "",
    "9. CONSEQUENCES OF VIOLATIONS",
    "- Minor violations may result in a warning or counselling",
    "- Serious violations may result in suspension or removal from the programme",
    "- Severe violations (e.g., harassment, fraud) may result in immediate dismissal",
    "",
    "10. REVIEW AND AMENDMENTS",
    "- This code of conduct will be reviewed annually",
    "- All participants will acknowledge receipt annually",
    "- Proposed amendments should be submitted to the programme coordinator",
    "",
    "ACKNOWLEDGEMENT: All participants must sign the code of conduct acknowledgement form.",
]
create_docx("data/raw/code_of_conduct.docx", "Code of Conduct", conduct_content)

# Expense Policy DOCX
expense_content = [
    "EXPENSE POLICY",
    "HCLTech Career Shaper Programme",
    "",
    "1. PURPOSE",
    "This policy establishes the guidelines for expense reimbursement for all participants in the HCLTech Career Shaper Programme. It aims to ensure transparency, consistency, and compliance with financial policies.",
    "",
    "2. REIMBURSABLE EXPENSES",
    "- Travel: Airfare, train tickets, mileage at the approved rate (₹15 per km)",
    "- Accommodation: Hotel rates up to ₹3,000 per night for out-of-town participation",
    "- Meals: Per diem rates as specified below",
    "- Training materials: Books, software, or equipment required for the programme",
    "- Examination fees: As approved by the programme coordinator",
    "",
    "3. MEAL PER DIEM RATES",
    "- Breakfast: ₹300 (if not provided)",
    "- Lunch: ₹500 (if not provided)",
    "- Dinner: ₹500 (if not provided)",
    "- Full day (all meals): ₹1,300 maximum",
    "",
    "4. EXPENSE SUBMISSION PROCEDURE",
    "- All expenses must be submitted within 30 days of incurrence",
    "- Receipts must be attached for all expenses exceeding ₹100",
    "- Missing receipts for amounts under ₹100 require a declaration form",
    "- Expense reports are submitted through the finance portal",
    "",
    "5. APPROVAL PROCESS",
    "- Department managers approve expense reports within their department",
    "- The programme coordinator has final approval authority",
    "- Approved expenses are reimbursed within 15 working days",
    "",
    "6. NON-REIMBURSABLE EXPENSES",
    "- Personal entertainment or luxury items",
    "- Alcohol (except during official business dinners with prior approval)",
    "- Fines, penalties, or late fees",
    "- expenses without proper documentation",
    
